import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook


SRC = Path(__file__).resolve().parents[1] / "src"
sys.path.insert(0, str(SRC))

from portable_agent.agent import TableAgent
from portable_agent.semantics import load_default_semantic_catalog
from portable_agent.sources import (
    SheetSelectionRequiredError,
    XlsxSourceError,
    XlsxTableSource,
    list_sheet_names,
    open_table_source,
)


class StaticGenerator:
    def __init__(self, payload):
        self.payload = payload

    def generate_plan(self, question, columns, semantic_definitions):
        return self.payload


class XlsxTableSourceTests(unittest.TestCase):
    def test_single_sheet_workbook_is_read_only_and_cited(self):
        with self.workbook_file() as path:
            before = path.read_bytes()
            source = open_table_source(path)
            agent = TableAgent(
                source,
                StaticGenerator({
                    "calculations": [{
                        "label": "Stunden",
                        "aggregation": "sum",
                        "column": "Einsatzstunden",
                    }]
                }),
                load_default_semantic_catalog(),
            )
            result = agent.ask("Einsatzstunden insgesamt")
            after = path.read_bytes()

        self.assertEqual(source.columns, ("Aktion", "Einsatzstunden", "Datum"))
        self.assertEqual(result.values, {"Stunden": 12.5})
        self.assertEqual([ref.row for ref in result.citations], [2, 4])
        self.assertEqual([ref.section for ref in result.citations], ["Daten", "Daten"])
        self.assertEqual(before, after)

    def test_multiple_sheets_require_explicit_selection(self):
        with self.workbook_file(multiple=True) as path:
            self.assertEqual(list_sheet_names(path), ("Daten", "Andere"))
            with self.assertRaisesRegex(
                SheetSelectionRequiredError,
                "explizite Auswahl erforderlich",
            ):
                XlsxTableSource(path)
            source = XlsxTableSource(path, sheet_name="Andere")
        self.assertEqual(source.sheet_name, "Andere")
        self.assertEqual(next(source.iter_rows()).values["Wert"], 9)

    def test_unknown_sheet_is_rejected(self):
        with self.workbook_file() as path:
            with self.assertRaisesRegex(XlsxSourceError, "nicht gefunden"):
                XlsxTableSource(path, sheet_name="Fehlt")

    def test_xlsm_container_is_read_without_macro_preservation(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "synthetic.xlsm"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Wert"])
            sheet.append([4])
            workbook.save(path)
            workbook.close()
            source = open_table_source(path)
        self.assertEqual(next(source.iter_rows()).values["Wert"], 4)

    def test_xlsm_latest_event_in_year_is_answered_deterministically(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "synthetic.xlsm"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Eingabe_Einsaetze"
            sheet.append(["Datum", "Aktion"])
            sheet.append([date(2025, 12, 31), "ALT"])
            sheet.append([date(2026, 2, 3), "FRUEH"])
            sheet.append([date(2026, 11, 19), "SPAET"])
            sheet.append([date(2027, 1, 1), "NEU"])
            workbook.save(path)
            workbook.close()
            before = path.read_bytes()
            source = open_table_source(path, sheet_name="Eingabe_Einsaetze")
            agent = TableAgent(
                source,
                StaticGenerator({}),
                load_default_semantic_catalog(),
            )

            result = agent.ask("Wann war der letzte Einsatz im Jahr 2026?")
            after = path.read_bytes()

        self.assertEqual(result.values, {"groups": [{
            "Aktion": "SPAET",
            "Letzter Einsatz": "2026-11-19",
        }]})
        self.assertEqual([ref.row for ref in result.citations], [4])
        self.assertTrue(all(ref.section == "Eingabe_Einsaetze" for ref in result.citations))
        self.assertEqual(before, after)

    def test_dates_durations_and_formulas_are_not_executed(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "types.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Typen"
            sheet.append(["Datum", "Dauer", "Formel"])
            sheet.append([date(2026, 9, 8), timedelta(hours=2, minutes=30), "=1+1"])
            workbook.save(path)
            workbook.close()
            source = XlsxTableSource(path)
            row = next(source.iter_rows()).values
        self.assertEqual(row["Datum"], "2026-09-08")
        self.assertEqual(row["Dauer"], 2.5)
        self.assertIsNone(row["Formel"])

    def test_duplicate_or_empty_headers_are_rejected(self):
        for headers, pattern in ((["A", "A"], "doppelte"), (["A", None, "C"], "leeren")):
            with self.subTest(headers=headers):
                with tempfile.TemporaryDirectory() as temporary:
                    path = Path(temporary) / "headers.xlsx"
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet.append(headers)
                    sheet.append([1] * len(headers))
                    workbook.save(path)
                    workbook.close()
                    with self.assertRaisesRegex(XlsxSourceError, pattern):
                        XlsxTableSource(path)

    def test_size_and_row_limits_are_enforced(self):
        with self.workbook_file() as path:
            with self.assertRaisesRegex(XlsxSourceError, "groesser als"):
                XlsxTableSource(path, max_bytes=2)
            with self.assertRaisesRegex(XlsxSourceError, "mehr als 1 Datenzeilen"):
                XlsxTableSource(path, max_rows=1)

    def workbook_file(self, multiple=False):
        class SyntheticWorkbook:
            def __enter__(inner_self):
                inner_self.temporary = tempfile.TemporaryDirectory()
                inner_self.path = Path(inner_self.temporary.name) / "synthetic.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Daten"
                sheet.append(["Aktion", "Einsatzstunden", "Datum"])
                sheet.append(["ALPHA", 7.5, date(2026, 9, 8)])
                sheet.append([None, None, None])
                sheet.append(["BETA", 5.0, date(2026, 9, 9)])
                if multiple:
                    other = workbook.create_sheet("Andere")
                    other.append(["Wert"])
                    other.append([9])
                workbook.save(inner_self.path)
                workbook.close()
                return inner_self.path

            def __exit__(inner_self, exc_type, exc_value, traceback):
                inner_self.temporary.cleanup()

        return SyntheticWorkbook()


if __name__ == "__main__":
    unittest.main()
