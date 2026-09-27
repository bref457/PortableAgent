"""Read-only XLSX/XLSM table source backed by openpyxl."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - explicit runtime guard below
    load_workbook = None

from .table import InMemoryTableSource


DEFAULT_MAX_BYTES = 100 * 1024 * 1024
DEFAULT_MAX_ROWS = 1_000_000
DEFAULT_MAX_COLUMNS = 500
_EXTENSIONS = frozenset({".xlsx", ".xlsm"})


class XlsxSourceError(ValueError):
    """An XLSX/XLSM source cannot be opened under the read-only policy."""


class XlsxDependencyError(XlsxSourceError):
    """The required local openpyxl package is unavailable."""


class SheetSelectionRequiredError(XlsxSourceError):
    """A workbook has multiple sheets and needs an explicit selection."""


class XlsxTableSource:
    """Snapshots one explicitly selected worksheet without modifying it."""

    def __init__(
        self,
        path: str | Path,
        *,
        sheet_name: str | None = None,
        header_row: int = 1,
        max_bytes: int = DEFAULT_MAX_BYTES,
        max_rows: int = DEFAULT_MAX_ROWS,
        max_columns: int = DEFAULT_MAX_COLUMNS,
    ) -> None:
        source_path = Path(path)
        _validate_request(source_path, header_row, max_bytes, max_rows, max_columns)
        _require_dependency()
        try:
            size = source_path.stat().st_size
        except OSError as exc:
            raise XlsxSourceError(f"Arbeitsmappe nicht lesbar: {exc}") from exc
        if size > max_bytes:
            raise XlsxSourceError(
                f"Arbeitsmappe ist groesser als das erlaubte Limit von {max_bytes} Bytes."
            )

        workbook = None
        try:
            workbook = load_workbook(
                filename=source_path,
                read_only=True,
                data_only=True,
                keep_vba=False,
                keep_links=False,
            )
            selected = _select_sheet(tuple(workbook.sheetnames), sheet_name)
            worksheet = workbook[selected]
            rows, row_numbers, headers = _read_sheet(
                worksheet,
                header_row=header_row,
                max_rows=max_rows,
                max_columns=max_columns,
            )
        except XlsxSourceError:
            raise
        except Exception as exc:
            raise XlsxSourceError(f"Arbeitsmappe konnte nicht gelesen werden: {exc}") from exc
        finally:
            if workbook is not None:
                workbook.close()

        self._columns = headers
        self._sheet_name = selected
        self._delegate = InMemoryTableSource(
            rows,
            source_id=f"xlsx:{source_path.name}:{selected}",
            display_name=source_path.name,
            section=selected,
            row_numbers=row_numbers,
        )

    @property
    def columns(self) -> tuple[str, ...]:
        return self._columns

    @property
    def sheet_name(self) -> str:
        return self._sheet_name

    def iter_rows(self):
        return self._delegate.iter_rows()


def list_sheet_names(path: str | Path) -> tuple[str, ...]:
    """List worksheet names without loading worksheet cell data."""
    source_path = Path(path)
    _validate_request(
        source_path,
        header_row=1,
        max_bytes=DEFAULT_MAX_BYTES,
        max_rows=DEFAULT_MAX_ROWS,
        max_columns=DEFAULT_MAX_COLUMNS,
    )
    _require_dependency()
    workbook = None
    try:
        workbook = load_workbook(
            filename=source_path,
            read_only=True,
            data_only=True,
            keep_vba=False,
            keep_links=False,
        )
        return tuple(workbook.sheetnames)
    except Exception as exc:
        raise XlsxSourceError(f"Arbeitsmappe konnte nicht gelesen werden: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _validate_request(path, header_row, max_bytes, max_rows, max_columns) -> None:
    if path.suffix.casefold() not in _EXTENSIONS:
        raise XlsxSourceError("Nur Dateien mit der Endung .xlsx oder .xlsm sind erlaubt.")
    if not path.exists() or not path.is_file():
        raise XlsxSourceError("Arbeitsmappe wurde nicht gefunden.")
    for value, name in (
        (header_row, "header_row"),
        (max_bytes, "max_bytes"),
        (max_rows, "max_rows"),
        (max_columns, "max_columns"),
    ):
        if type(value) is not int or value <= 0:
            raise XlsxSourceError(f"{name} muss eine positive ganze Zahl sein.")


def _require_dependency() -> None:
    if load_workbook is None:
        raise XlsxDependencyError(
            "openpyxl ist lokal nicht installiert; XLSX/XLSM kann nicht gelesen werden."
        )


def _select_sheet(sheet_names: tuple[str, ...], requested: str | None) -> str:
    if not sheet_names:
        raise XlsxSourceError("Arbeitsmappe enthaelt keine Tabellenblaetter.")
    if requested is None:
        if len(sheet_names) == 1:
            return sheet_names[0]
        raise SheetSelectionRequiredError(
            "Arbeitsmappe hat mehrere Tabellenblaetter; explizite Auswahl erforderlich: "
            + ", ".join(sheet_names)
        )
    if requested not in sheet_names:
        raise XlsxSourceError(
            f"Tabellenblatt '{requested}' wurde nicht gefunden. Verfuegbar: "
            + ", ".join(sheet_names)
        )
    return requested


def _read_sheet(worksheet, *, header_row, max_rows, max_columns):
    iterator = worksheet.iter_rows(min_row=header_row, values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration as exc:
        raise XlsxSourceError("Tabellenblatt enthaelt keine Kopfzeile.") from exc
    if len(raw_headers) > max_columns:
        raise XlsxSourceError(f"Tabellenblatt hat mehr als {max_columns} Spalten.")
    headers = tuple(
        str(value).strip() if value is not None else ""
        for value in raw_headers
    )
    while headers and not headers[-1]:
        headers = headers[:-1]
    if not headers or any(not header for header in headers):
        raise XlsxSourceError("Kopfzeile enthaelt einen leeren Spaltennamen.")
    if len(set(headers)) != len(headers):
        raise XlsxSourceError("Kopfzeile enthaelt doppelte Spaltennamen.")

    rows: list[dict[str, object]] = []
    row_numbers: list[int] = []
    for excel_row, raw_row in enumerate(iterator, start=header_row + 1):
        values = tuple(raw_row[: len(headers)])
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(rows) >= max_rows:
            raise XlsxSourceError(
                f"Tabellenblatt enthaelt mehr als {max_rows} Datenzeilen."
            )
        rows.append({
            header: _normalize_cell(value)
            for header, value in zip(headers, values, strict=True)
        })
        row_numbers.append(excel_row)
    return rows, row_numbers, headers


def _normalize_cell(value: Any) -> object:
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds() / 3600
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value
